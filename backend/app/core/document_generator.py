"""
Génération de documents Word (.docx) pour les contrats de travail.
Utilisé à la fois par l'espace RH (Phase 2C) et l'espace employé (Phase 2D).
"""
import io
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

from app.core.ai_service import parse_structured_content

ENTREPRISE_NOM = "Enterprise RH"


def _format_montant(montant: int) -> str:
    """Formatte un montant entier avec des espaces comme séparateur de milliers."""
    return f"{montant:,}".replace(",", " ") + " DT"


def _format_date(d) -> str:
    return d.strftime("%d/%m/%Y") if d else "-"


def _render_article_content(doc: Document, content_text: str):
    """
    Rend le contenu d'un article dans le document Word.
    Prend en charge à la fois le texte brut classique et le JSON structuré
    (paragraphes, tableaux Word avec en-têtes et bordures, ou mixte).
    """
    if not content_text or not content_text.strip():
        doc.add_paragraph().add_run("(Contenu non renseigné)").italic = True
        return

    structured = parse_structured_content(content_text)
    if not structured:
        # Contenu texte brut standard (rétrocompatibilité totale)
        doc.add_paragraph(content_text)
        return

    # Contenu structuré avec blocs
    blocks = structured.get("blocks", [])
    for block in blocks:
        btype = block.get("type", "paragraph")
        if btype == "paragraph":
            text = block.get("content", "").strip()
            if text:
                doc.add_paragraph(text)
        elif btype == "table":
            headers = block.get("headers", [])
            rows = block.get("rows", [])
            if headers:
                col_count = len(headers)
                t = doc.add_table(rows=0, cols=col_count)
                t.style = "Light Grid Accent 1"

                # Ligne d'en-tête
                hdr_cells = t.add_row().cells
                for idx, h in enumerate(headers):
                    hdr_cells[idx].text = str(h)
                    if hdr_cells[idx].paragraphs[0].runs:
                        hdr_cells[idx].paragraphs[0].runs[0].bold = True

                # Lignes de données
                for row_data in rows:
                    row_cells = t.add_row().cells
                    for idx, val in enumerate(row_data):
                        if idx < col_count:
                            row_cells[idx].text = str(val) if val is not None else ""

                # Espace après le tableau
                doc.add_paragraph()



def generer_contrat_word(contrat) -> io.BytesIO:
    """
    Génère un document Word professionnel pour un contrat donné.

    :param contrat: instance ORM `Contrat` avec les relations `employe` et
                     `articles` déjà chargées (joinedload recommandé).
    :return: buffer BytesIO positionné au début, prêt à être streamé.
    """
    employe = contrat.employe
    articles = contrat.articles or []

    doc = Document()

    # Marges document
    section = doc.sections[0]
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.5)

    # ─── Titre principal ────────────────────────────────────────────
    titre = doc.add_heading("CONTRAT DE TRAVAIL", level=0)
    titre.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in titre.runs:
        run.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)

    ref_para = doc.add_paragraph()
    ref_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    ref_run = ref_para.add_run(f"Référence : {contrat.reference}")
    ref_run.italic = True
    ref_run.font.size = Pt(11)
    ref_run.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)

    doc.add_paragraph()

    # ─── Préambule / parties ────────────────────────────────────────
    p = doc.add_paragraph()
    p.add_run("Entre les soussignés :").bold = True

    doc.add_paragraph()

    p_employeur = doc.add_paragraph()
    p_employeur.add_run(ENTREPRISE_NOM).bold = True
    p_employeur.add_run(", représentée par sa Direction des Ressources Humaines,")
    p_employeur.add_run("\nCi-après désignée « l'Employeur »,")
    doc.add_paragraph("D'une part,")

    doc.add_paragraph()
    doc.add_paragraph().add_run("ET").bold = True
    doc.add_paragraph()

    p_salarie = doc.add_paragraph()
    if employe:
        p_salarie.add_run(f"{employe.prenom} {employe.nom}").bold = True
        p_salarie.add_run(f"\nMatricule : {employe.matricule}")
        p_salarie.add_run(f"\nPoste : {employe.poste or 'Non défini'}")
        if employe.departement:
            p_salarie.add_run(f"\nDépartement : {employe.departement}")
    else:
        p_salarie.add_run("Employé non renseigné").italic = True
    p_salarie.add_run("\nCi-après désigné(e) « le Salarié »,")
    doc.add_paragraph("D'autre part,")

    doc.add_paragraph()
    doc.add_paragraph("Il a été convenu et arrêté ce qui suit :")
    doc.add_paragraph()

    # ─── Détails du contrat ─────────────────────────────────────────
    doc.add_heading("Article préliminaire — Conditions générales", level=1)

    table = doc.add_table(rows=0, cols=2)
    table.style = "Light Grid Accent 1"

    def add_row(label: str, value: str):
        row = table.add_row()
        row.cells[0].text = label
        row.cells[1].text = value
        row.cells[0].paragraphs[0].runs[0].bold = True

    add_row("Référence du contrat", getattr(contrat, "reference", "-"))
    tc = getattr(contrat, "type_contrat", None)
    if tc:
        add_row("Type de contrat", str(tc))
    add_row("Date de début", _format_date(contrat.date_debut))
    add_row("Date de fin", _format_date(contrat.date_fin) if contrat.date_fin else "Durée indéterminée")
    add_row("Salaire mensuel brut", _format_montant(contrat.salaire_mensuel))

    doc.add_paragraph()

    # ─── Articles / clauses ─────────────────────────────────────────
    if articles:
        doc.add_heading("Clauses contractuelles", level=1)
        for i, article in enumerate(articles, start=1):
            doc.add_heading(f"Article {i} — {article.titre}", level=2)
            _render_article_content(doc, article.contenu_par_defaut)

    doc.add_paragraph()
    doc.add_paragraph()

    # ─── Signatures ─────────────────────────────────────────────────
    doc.add_paragraph("Fait en deux exemplaires originaux.")
    doc.add_paragraph()

    sig_table = doc.add_table(rows=2, cols=2)
    sig_table.cell(0, 0).text = "Signature de l'employeur"
    sig_table.cell(0, 1).text = "Signature de l'employé"
    sig_table.cell(1, 0).text = "\n\n\n_______________________"
    sig_table.cell(1, 1).text = "\n\n\n_______________________"

    for row in sig_table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                if paragraph.runs:
                    paragraph.runs[0].bold = True

    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer
