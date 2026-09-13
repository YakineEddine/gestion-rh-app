import os
import re
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, color_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=140, bottom=140, left=180, right=180):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'''
        <w:tcMar {nsdecls("w")}>
            <w:top w:w="{top}" w:type="dxa"/>
            <w:bottom w:w="{bottom}" w:type="dxa"/>
            <w:left w:w="{left}" w:type="dxa"/>
            <w:right w:w="{right}" w:type="dxa"/>
        </w:tcMar>
    ''')
    tcPr.append(tcMar)

def set_cell_left_border(cell, color_hex="C00000", sz="36"):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(f'''
        <w:tcBorders {nsdecls("w")}>
            <w:top w:val="none"/>
            <w:left w:val="single" w:sz="{sz}" w:space="0" w:color="{color_hex}"/>
            <w:bottom w:val="none"/>
            <w:right w:val="none"/>
        </w:tcBorders>
    ''')
    tcPr.append(borders)

def set_table_borders(table, color_hex="CBD5E1"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(f'''
        <w:tblBorders {nsdecls("w")}>
            <w:top w:val="single" w:sz="6" w:space="0" w:color="{color_hex}"/>
            <w:left w:val="single" w:sz="6" w:space="0" w:color="{color_hex}"/>
            <w:bottom w:val="single" w:sz="6" w:space="0" w:color="{color_hex}"/>
            <w:right w:val="single" w:sz="6" w:space="0" w:color="{color_hex}"/>
            <w:insideH w:val="single" w:sz="4" w:space="0" w:color="{color_hex}"/>
            <w:insideV w:val="single" w:sz="4" w:space="0" w:color="{color_hex}"/>
        </w:tblBorders>
    ''')
    tblPr.append(borders)

def add_callout(doc, text, color_hex="C00000", bg_hex="F8FAFC"):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Cm(16.0)
    
    cell = table.cell(0, 0)
    set_cell_background(cell, bg_hex)
    set_cell_margins(cell, top=160, bottom=160, left=240, right=200)
    set_cell_left_border(cell, color_hex=color_hex, sz="32")
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run(text)
    run.font.name = "Calibri"
    run.font.size = Pt(10)
    run.font.italic = True
    run.font.color.rgb = RGBColor(71, 85, 105)
    
    # Empty line after callout
    p_after = doc.add_paragraph()
    p_after.paragraph_format.space_before = Pt(0)
    p_after.paragraph_format.space_after = Pt(6)

def add_code_block(doc, code_text):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Cm(16.0)
    
    cell = table.cell(0, 0)
    set_cell_background(cell, "F1F5F9")
    set_cell_margins(cell, top=140, bottom=140, left=180, right=180)
    
    # Border
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(f'''
        <w:tcBorders {nsdecls("w")}>
            <w:top w:val="single" w:sz="6" w:space="0" w:color="CBD5E1"/>
            <w:left w:val="single" w:sz="6" w:space="0" w:color="CBD5E1"/>
            <w:bottom w:val="single" w:sz="6" w:space="0" w:color="CBD5E1"/>
            <w:right w:val="single" w:sz="6" w:space="0" w:color="CBD5E1"/>
        </w:tcBorders>
    ''')
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.0
    run = p.add_run(code_text.strip())
    run.font.name = "Consolas"
    run.font.size = Pt(8.5)
    run.font.color.rgb = RGBColor(30, 41, 59)
    
    p_after = doc.add_paragraph()
    p_after.paragraph_format.space_before = Pt(0)
    p_after.paragraph_format.space_after = Pt(6)

def build_word_document(md_path, cover_img_path, output_docx_path):
    doc = Document()
    
    # --- SECTION 1: COVER PAGE (Zero Margins) ---
    sec1 = doc.sections[0]
    sec1.top_margin = Cm(0)
    sec1.bottom_margin = Cm(0)
    sec1.left_margin = Cm(0)
    sec1.right_margin = Cm(0)
    sec1.page_width = Cm(21.0)
    sec1.page_height = Cm(29.7)
    
    p_cov = sec1.header.paragraphs[0]
    p_cov.text = "" # Clean header
    
    # Insert Full Page Cover Image
    p_cover = doc.add_paragraph()
    p_cover.paragraph_format.space_before = Pt(0)
    p_cover.paragraph_format.space_after = Pt(0)
    run_cover = p_cover.add_run()
    run_cover.add_picture(cover_img_path, width=Cm(21.0), height=Cm(29.7))
    
    # --- SECTION 2: BODY (Standard Margins) ---
    sec2 = doc.add_section(docx.enum.section.WD_SECTION.NEW_PAGE)
    sec2.top_margin = Cm(2.5)
    sec2.bottom_margin = Cm(2.5)
    sec2.left_margin = Cm(2.5)
    sec2.right_margin = Cm(2.5)
    sec2.page_width = Cm(21.0)
    sec2.page_height = Cm(29.7)
    
    # Unlink header/footer from cover
    sec2.header.is_linked_to_previous = False
    sec2.footer.is_linked_to_previous = False
    
    # Header
    h_para = sec2.header.paragraphs[0]
    h_para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    h_run = h_para.add_run("Rapport de Stage d'Immersion en Entreprise | CSI Digital – ESPRIT")
    h_run.font.name = "Calibri"
    h_run.font.size = Pt(8.5)
    h_run.font.color.rgb = RGBColor(148, 163, 184)
    
    # Footer
    f_para = sec2.footer.paragraphs[0]
    f_para.alignment = WD_ALIGN_PARAGRAPH.LEFT
    f_run = f_para.add_run("SAHLI Yakine Eddine — Génie Informatique (2025/2026)")
    f_run.font.name = "Calibri"
    f_run.font.size = Pt(8.5)
    f_run.font.color.rgb = RGBColor(148, 163, 184)
    
    # Page Number in footer (right)
    # In python-docx, add page field
    f_p2 = sec2.footer.add_paragraph()
    f_p2.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    f_run2 = f_p2.add_run("Page ")
    f_run2.font.name = "Calibri"
    f_run2.font.size = Pt(8.5)
    f_run2.font.color.rgb = RGBColor(148, 163, 184)
    f_xml = parse_xml(r'<w:fldSimple %s w:instr="PAGE"/>' % nsdecls('w'))
    f_p2._p.append(f_xml)

    # Read markdown content
    with open(md_path, "r", encoding="utf-8") as f:
        md_text = f.read()

    # Skip the Markdown cover page section since we inserted the full HD cover
    # Look for "TABLE DES MATIÈRES"
    idx_toc = md_text.find("# TABLE DES MATIÈRES")
    if idx_toc != -1:
        content = md_text[idx_toc:]
    else:
        content = md_text

    # Parse content line by line with states
    lines = content.splitlines()
    i = 0
    in_code_block = False
    code_lines = []
    in_table = False
    table_lines = []

    while i < len(lines):
        line = lines[i]

        # Handle code blocks (```)
        if line.strip().startswith("```"):
            if in_code_block:
                in_code_block = False
                add_code_block(doc, "\n".join(code_lines))
                code_lines = []
            else:
                in_code_block = True
                code_lines = []
            i += 1
            continue

        if in_code_block:
            code_lines.append(line)
            i += 1
            continue

        # Handle tables (| ... |)
        if line.strip().startswith("|") and line.strip().endswith("|"):
            in_table = True
            table_lines.append(line)
            i += 1
            continue
        elif in_table:
            # End of table
            in_table = False
            render_markdown_table(doc, table_lines)
            table_lines = []
            # do not continue, process current line

        # Horizontal rules
        if line.strip() in ["---", "--- ---", "═══════════════════════════════════════════════════════════════"]:
            i += 1
            continue

        # Callouts (> ...)
        if line.strip().startswith(">"):
            callout_text = re.sub(r"^>\s*", "", line.strip())
            add_callout(doc, callout_text)
            i += 1
            continue

        # Headings
        if line.startswith("# "):
            h_text = line[2:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(20)
            p.paragraph_format.space_after = Pt(10)
            p.paragraph_format.keep_with_next = True
            if "Chapitre" in h_text or "Introduction" in h_text or "Conclusion" in h_text or "TABLE DES" in h_text:
                p.paragraph_format.page_break_before = True
            
            run = p.add_run(h_text)
            run.font.name = "Calibri"
            run.font.size = Pt(18)
            run.font.bold = True
            run.font.color.rgb = RGBColor(192, 0, 0) # ESPRIT Red
            i += 1
            continue

        if line.startswith("## "):
            h_text = line[3:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(h_text)
            run.font.name = "Calibri"
            run.font.size = Pt(14)
            run.font.bold = True
            run.font.color.rgb = RGBColor(30, 41, 59) # Slate
            i += 1
            continue

        if line.startswith("### "):
            h_text = line[4:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(h_text)
            run.font.name = "Calibri"
            run.font.size = Pt(12)
            run.font.bold = True
            run.font.color.rgb = RGBColor(71, 85, 105)
            i += 1
            continue

        # Bullet list items (- or *)
        if line.strip().startswith("- ") or line.strip().startswith("* "):
            bullet_text = line.strip()[2:].strip()
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.15
            render_inline_formatting(p, bullet_text)
            i += 1
            continue

        # Numbered lists (1. 2. etc.)
        num_match = re.match(r"^(\d+)\.\s+(.*)$", line.strip())
        if num_match:
            num_text = num_match.group(2)
            p = doc.add_paragraph(style='List Number')
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.15
            render_inline_formatting(p, num_text)
            i += 1
            continue

        # Empty lines
        if not line.strip():
            i += 1
            continue

        # Standard Paragraph
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.line_spacing = 1.15
        render_inline_formatting(p, line.strip())
        i += 1

    # End of document check
    if in_table and table_lines:
        render_markdown_table(doc, table_lines)

    doc.save(output_docx_path)
    print(f"Document Word successfully generated: {output_docx_path}")

def render_markdown_table(doc, lines):
    # Filter header separator line (|---|---|)
    filtered = []
    for l in lines:
        if re.search(r"\|(?:\s*:?-+:?\s*\|)+", l):
            continue
        # Split by |
        parts = [c.strip() for c in l.split("|")[1:-1]]
        if parts:
            filtered.append(parts)

    if not filtered:
        return

    num_cols = max(len(row) for row in filtered)
    num_rows = len(filtered)
    
    table = doc.add_table(rows=num_rows, cols=num_cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    set_table_borders(table, "CBD5E1")
    
    for r_idx, row in enumerate(filtered):
        is_header = (r_idx == 0)
        for c_idx in range(num_cols):
            cell = table.cell(r_idx, c_idx)
            val = row[c_idx] if c_idx < len(row) else ""
            
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.05
            
            if is_header:
                set_cell_background(cell, "1E293B") # Dark Slate Header
                set_cell_margins(cell, top=140, bottom=140, left=140, right=140)
                run = p.add_run(val)
                run.font.name = "Calibri"
                run.font.size = Pt(9.5)
                run.font.bold = True
                run.font.color.rgb = RGBColor(255, 255, 255)
            else:
                bg = "F8FAFC" if (r_idx % 2 == 1) else "FFFFFF"
                set_cell_background(cell, bg)
                set_cell_margins(cell, top=100, bottom=100, left=120, right=120)
                render_inline_formatting(p, val, font_size=9.5)
                
    p_after = doc.add_paragraph()
    p_after.paragraph_format.space_before = Pt(0)
    p_after.paragraph_format.space_after = Pt(6)

def render_inline_formatting(paragraph, text, font_size=11):
    # Regex parse bold (**text**), code (`text`), italic (*text*)
    tokens = re.split(r"(\*\*.*?\*\*|`.*?`|\*.*?\*)", text)
    for token in tokens:
        if not token:
            continue
        if token.startswith("**") and token.endswith("**"):
            run = paragraph.add_run(token[2:-2])
            run.font.name = "Calibri"
            run.font.size = Pt(font_size)
            run.font.bold = True
            run.font.color.rgb = RGBColor(30, 41, 59)
        elif token.startswith("`") and token.endswith("`"):
            run = paragraph.add_run(token[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(font_size - 1)
            run.font.color.rgb = RGBColor(180, 20, 20)
        elif token.startswith("*") and token.endswith("*"):
            run = paragraph.add_run(token[1:-1])
            run.font.name = "Calibri"
            run.font.size = Pt(font_size)
            run.font.italic = True
            run.font.color.rgb = RGBColor(71, 85, 105)
        else:
            run = paragraph.add_run(token)
            run.font.name = "Calibri"
            run.font.size = Pt(font_size)
            run.font.color.rgb = RGBColor(30, 41, 59)

if __name__ == "__main__":
    md_file = r"C:\Users\Home\.gemini\antigravity-ide\brain\0cf11b19-f886-41ea-9bfe-e1f777e3431a\rapport_stage.md"
    cover_img = r"C:\Users\Home\.gemini\antigravity-ide\scratch\gestion-rh-app\cover_esprit_hd.png"
    out_docx = r"C:\Users\Home\.gemini\antigravity-ide\scratch\gestion-rh-app\Rapport_Stage_Yakine_Eddine_ESPRIT.docx"
    build_word_document(md_file, cover_img, out_docx)
