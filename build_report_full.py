# -*- coding: utf-8 -*-
"""
Générateur Complet du Rapport de Stage Académique et Professionnel
Auteur : Sahli Yakine Eddine
Encadrante : Mme Kharbech Rayen
Établissement : ESPRIT
Entreprise : CSI Digital
Sujet : "Conception et développement d'une application web RH permettant de gérer les employés et les contrats de travail."
"""

import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls

BASE_DIR = r"c:\Users\Home\.gemini\antigravity-ide\scratch\gestion-rh-app"
ASSETS_DIR = os.path.join(BASE_DIR, "report_assets")
OUTPUT_PATH = os.path.join(BASE_DIR, "Rapport_de_Stage_Sahli_Yakine_Eddine_CSI_Digital.docx")

# Palette Corporate / Académique
COLOR_PRIMARY = RGBColor(27, 54, 93)     # #1B365D (Deep Navy)
COLOR_SECONDARY = RGBColor(43, 108, 176) # #2B6CB0 (Slate Blue)
COLOR_TEXT = RGBColor(45, 55, 72)        # #2D3748 (Charcoal Text)
COLOR_MUTED = RGBColor(113, 128, 150)    # #718096 (Slate Muted)
COLOR_ACCENT = RGBColor(44, 122, 123)    # #2C7A7B (Teal Accent)

HEX_PRIMARY = "1B365D"
HEX_SECONDARY = "2B6CB0"
HEX_LIGHT_BG = "F0F4F8"
HEX_ZEBRA = "F8FAFC"
HEX_BORDER = "CBD5E0"
HEX_CALLOUT = "EBF8FF"

def set_cell_background(cell, fill_hex):
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def set_cell_borders(cell, top="none", bottom="none", left="none", right="none", color="CBD5E0", sz="4"):
    tcPr = cell._element.get_or_add_tcPr()
    borders_xml = f'<w:tcBorders {nsdecls("w")}>'
    for side, val in [("top", top), ("bottom", bottom), ("left", left), ("right", right)]:
        if val == "single":
            borders_xml += f'<w:{side} w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        elif val == "none":
            borders_xml += f'<w:{side} w:val="none"/>'
    borders_xml += '</w:tcBorders>'
    tcPr.append(parse_xml(borders_xml))

def add_page_number(run):
    fldChar1 = parse_xml(r'<w:fldChar %s w:fldCharType="begin"/>' % nsdecls('w'))
    instrText = parse_xml(r'<w:instrText %s xml:space="preserve"> PAGE </w:instrText>' % nsdecls('w'))
    fldChar2 = parse_xml(r'<w:fldChar %s w:fldCharType="separate"/>' % nsdecls('w'))
    fldChar3 = parse_xml(r'<w:fldChar %s w:fldCharType="end"/>' % nsdecls('w'))
    r = run._r
    r.append(fldChar1)
    r.append(instrText)
    r.append(fldChar2)
    r.append(fldChar3)

def add_total_pages(run):
    fldChar1 = parse_xml(r'<w:fldChar %s w:fldCharType="begin"/>' % nsdecls('w'))
    instrText = parse_xml(r'<w:instrText %s xml:space="preserve"> NUMPAGES </w:instrText>' % nsdecls('w'))
    fldChar2 = parse_xml(r'<w:fldChar %s w:fldCharType="separate"/>' % nsdecls('w'))
    fldChar3 = parse_xml(r'<w:fldChar %s w:fldCharType="end"/>' % nsdecls('w'))
    r = run._r
    r.append(fldChar1)
    r.append(instrText)
    r.append(fldChar2)
    r.append(fldChar3)

def add_heading_1(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(16)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.name = "Calibri"
    run.font.size = Pt(16)
    run.font.bold = True
    run.font.color.rgb = COLOR_PRIMARY
    return p

def add_heading_2(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.name = "Calibri"
    run.font.size = Pt(13)
    run.font.bold = True
    run.font.color.rgb = COLOR_SECONDARY
    return p

def add_heading_3(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(9)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.name = "Calibri"
    run.font.size = Pt(11)
    run.font.bold = True
    run.font.color.rgb = COLOR_ACCENT
    return p

def add_paragraph(doc, text, bold_prefix=None, space_after=5, italic=False):
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_after = Pt(space_after)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if bold_prefix:
        r_bold = p.add_run(bold_prefix + " ")
        r_bold.font.name = "Calibri"
        r_bold.font.size = Pt(10)
        r_bold.font.bold = True
        r_bold.font.color.rgb = COLOR_PRIMARY
    run = p.add_run(text)
    run.font.name = "Calibri"
    run.font.size = Pt(10)
    run.font.color.rgb = COLOR_TEXT
    if italic:
        run.font.italic = True
    return p

def add_bullet(doc, text, bold_prefix=None):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_after = Pt(2.5)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if bold_prefix:
        r_bold = p.add_run(bold_prefix + " : ")
        r_bold.font.name = "Calibri"
        r_bold.font.size = Pt(9.5)
        r_bold.font.bold = True
        r_bold.font.color.rgb = COLOR_PRIMARY
    run = p.add_run(text)
    run.font.name = "Calibri"
    run.font.size = Pt(9.5)
    run.font.color.rgb = COLOR_TEXT
    return p

def add_callout(doc, text, title="Remarque technique"):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    set_cell_background(cell, HEX_CALLOUT)
    set_cell_margins(cell, top=90, bottom=90, left=160, right=160)
    set_cell_borders(cell, left="single", color=HEX_SECONDARY, sz="24")
    
    p = cell.paragraphs[0]
    p.paragraph_format.line_spacing = 1.12
    p.paragraph_format.space_after = Pt(0)
    
    r_title = p.add_run(f"📌 {title} : ")
    r_title.font.name = "Calibri"
    r_title.font.bold = True
    r_title.font.size = Pt(9.5)
    r_title.font.color.rgb = COLOR_SECONDARY
    
    r_body = p.add_run(text)
    r_body.font.name = "Calibri"
    r_body.font.italic = True
    r_body.font.size = Pt(9.5)
    r_body.font.color.rgb = COLOR_TEXT

    sp = doc.add_paragraph()
    sp.paragraph_format.space_after = Pt(4)

def add_styled_table(doc, headers, rows_data, col_widths=None):
    tbl = doc.add_table(rows=len(rows_data) + 1, cols=len(headers))
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False

    # Header row
    hdr_cells = tbl.rows[0].cells
    for i, h in enumerate(headers):
        hdr_cells[i].text = h
        set_cell_background(hdr_cells[i], HEX_PRIMARY)
        set_cell_margins(hdr_cells[i], top=90, bottom=90, left=110, right=110)
        set_cell_borders(hdr_cells[i], top="single", bottom="single", left="single", right="single", color=HEX_PRIMARY, sz="4")
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for r in p.runs:
            r.font.name = "Calibri"
            r.font.bold = True
            r.font.size = Pt(9)
            r.font.color.rgb = RGBColor(255, 255, 255)

    # Data rows
    for r_idx, row in enumerate(rows_data):
        row_cells = tbl.rows[r_idx + 1].cells
        bg_color = HEX_ZEBRA if r_idx % 2 == 1 else "FFFFFF"
        for c_idx, val in enumerate(row):
            row_cells[c_idx].text = str(val)
            set_cell_background(row_cells[c_idx], bg_color)
            set_cell_margins(row_cells[c_idx], top=65, bottom=65, left=95, right=95)
            set_cell_borders(row_cells[c_idx], top="single", bottom="single", left="single", right="single", color=HEX_BORDER, sz="4")
            p = row_cells[c_idx].paragraphs[0]
            p.paragraph_format.line_spacing = 1.1
            p.paragraph_format.space_after = Pt(1)
            if c_idx == 0 and len(str(val)) < 15:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            for r in p.runs:
                r.font.name = "Calibri"
                r.font.size = Pt(8.5)
                r.font.color.rgb = COLOR_TEXT

    if col_widths:
        for row in tbl.rows:
            for idx, width in enumerate(col_widths):
                row.cells[idx].width = Inches(width)

    sp = doc.add_paragraph()
    sp.paragraph_format.space_after = Pt(5)
    return tbl

def add_figure(doc, img_name, caption, width=5.5):
    img_path = os.path.join(ASSETS_DIR, img_name)
    if not os.path.exists(img_path):
        print(f"Warning: Figure not found: {img_path}")
        return
    
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    
    run = p.add_run()
    run.add_picture(img_path, width=Inches(width))
    
    cp = doc.add_paragraph()
    cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cp.paragraph_format.space_before = Pt(0)
    cp.paragraph_format.space_after = Pt(7)
    
    parts = caption.split(" — ", 1)
    c_bold = cp.add_run(parts[0] + " — ")
    c_bold.font.name = "Calibri"
    c_bold.font.bold = True
    c_bold.font.size = Pt(8.5)
    c_bold.font.color.rgb = COLOR_PRIMARY
    
    if len(parts) > 1:
        c_text = cp.add_run(parts[1])
        c_text.font.name = "Calibri"
        c_text.font.italic = True
        c_text.font.size = Pt(8.5)
        c_text.font.color.rgb = COLOR_MUTED

print("Setup completed. Building document...")
